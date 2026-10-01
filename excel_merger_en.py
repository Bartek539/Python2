"""
Multi-Sheet Excel Merger Studio  (v2, English UI)

Merges .xlsx/.xlsm/.xlsb/.xls/.csv files from one folder into a single workbook,
sheet by sheet (sheets with the same name are stacked on top of each other).

Required: pandas, openpyxl  (optional: xlrd for .xls, pyxlsb for .xlsb)
"""
from __future__ import annotations

import csv
import ctypes
import json
import logging
import os
import queue
import re
import sys
import threading
import traceback
from dataclasses import dataclass, field
from datetime import datetime
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

APP_NAME = "ExcelMergerStudio"


# =============================================================================
# 1. PATHS, LOGGING, SYSTEM ERRORS
# =============================================================================
def get_base_dir() -> Path:
    """Directory of the .exe (PyInstaller) or of the .py script."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def get_data_dir() -> Path:
    """Writable directory for app data (log, settings).

    Uses %LOCALAPPDATA% instead of the .exe folder, which can be read-only
    (e.g. Program Files) - the log would then silently stop working.
    """
    root = os.environ.get("LOCALAPPDATA")
    candidate = Path(root) / APP_NAME if root else get_base_dir()
    try:
        candidate.mkdir(parents=True, exist_ok=True)
        return candidate
    except OSError:
        return get_base_dir()


DATA_DIR = get_data_dir()
LOG_FILE = DATA_DIR / "merger_system.log"
SETTINGS_FILE = DATA_DIR / "settings.json"


def setup_logging() -> logging.Logger:
    logger = logging.getLogger(APP_NAME)
    logger.setLevel(logging.INFO)
    if not logger.handlers:
        try:
            handler = RotatingFileHandler(
                LOG_FILE, maxBytes=1_000_000, backupCount=3, encoding="utf-8"
            )
            handler.setFormatter(
                logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
            )
            logger.addHandler(handler)
        except OSError:
            logger.addHandler(logging.NullHandler())
    return logger


log = setup_logging()


def show_native_error(title: str, message: str) -> None:
    """WinAPI error box - works even if Tkinter fails to load."""
    try:
        ctypes.windll.user32.MessageBoxW(0, message, title, 0x10 | 0x10000)
    except Exception:
        print(f"{title}: {message}", file=sys.stderr)


def enable_dpi_awareness() -> None:
    """Sharp text on HiDPI screens (must be called before creating Tk)."""
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass


log.info("=== APPLICATION INITIALIZATION STARTED ===")

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    import pandas as pd
    import openpyxl  # noqa: F401  (required by pandas for .xlsx; importing = fail early)

    log.info("Core libraries imported successfully.")
except Exception as exc:  # noqa: BLE001
    log.error("Import failed:\n%s", traceback.format_exc())
    show_native_error(
        "Startup Error",
        f"Failed to load required libraries:\n\n{exc}\n\nDetails: {LOG_FILE}",
    )
    sys.exit(1)


# =============================================================================
# 2. MERGE LOGIC (no GUI dependencies - easy to test)
# =============================================================================
SUPPORTED_EXTS = {".xlsx", ".xlsm", ".xlsb", ".xls", ".csv"}
OUTPUT_PREFIX = "merged_output_"
SOURCE_COL = "_Source_File"
EXCEL_MAX_DATA_ROWS = 1_048_575  # Excel row limit minus the header row
CSV_ENCODINGS = ("utf-8-sig", "cp1250", "latin-1")

_BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class MergeOptions:
    input_dir: Path
    output_dir: Path
    order_file: Optional[Path] = None
    skip_top: int = 0
    skip_bottom: int = 0
    add_source_column: bool = False


@dataclass
class MergeResult:
    output_file: Optional[Path] = None
    files_total: int = 0
    files_ok: int = 0
    sheets: int = 0
    rows: int = 0
    cancelled: bool = False
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


ProgressCb = Callable[[int, int, str], None]


def excel_engine(ext: str) -> str:
    if ext in {".xlsx", ".xlsm"}:
        return "openpyxl"
    if ext == ".xlsb":
        return "pyxlsb"
    return "xlrd"


def natural_key(text: str) -> list:
    """Natural sort: file2 < file10."""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", text.casefold())]


def read_csv_robust(path: Path, **kwargs) -> "pd.DataFrame":
    """CSV with delimiter and encoding auto-detection (UTF-8 -> cp1250 -> latin-1)."""
    for enc in CSV_ENCODINGS:
        try:
            try:
                return pd.read_csv(path, sep=None, engine="python", encoding=enc, **kwargs)
            except (csv.Error, pd.errors.ParserError):
                # e.g. single-column file - the sniffer cannot detect a delimiter
                return pd.read_csv(path, sep=",", engine="python", encoding=enc, **kwargs)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not detect the text encoding of {path.name}")


def read_order_list(path: Path) -> List[str]:
    """First column of the order file -> list of codes (as text)."""
    ext = path.suffix.lower()
    if ext in {".csv", ".txt"}:
        text = ""
        for enc in CSV_ENCODINGS:
            try:
                text = path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        codes = [re.split(r"[;\t,]", line)[0] for line in text.splitlines()]
    else:
        df = pd.read_excel(path, usecols=[0], header=None, dtype=str, engine=excel_engine(ext))
        codes = df.iloc[:, 0].dropna().tolist()
    cleaned = [str(c).strip().strip('"').strip() for c in codes]
    return [c for c in cleaned if c]


def sort_files(
    files: List[Path], codes: List[str]
) -> Tuple[List[Path], List[str], List[Path]]:
    """Sorts files by the list of codes.

    Matching is case-insensitive; when several codes match, the longest wins
    (code "12" will not capture the file for code "123").
    Returns: (sorted files, codes without a file, files without a code).
    """
    lowered = [(i, c.casefold()) for i, c in enumerate(codes)]
    used_codes: set = set()
    ordered: List[Tuple[int, list, Path]] = []
    rest: List[Path] = []

    for p in files:
        name = p.name.casefold()
        hits = [(i, c) for i, c in lowered if c in name]
        if hits:
            best_idx, _ = max(hits, key=lambda h: (len(h[1]), -h[0]))
            used_codes.add(best_idx)
            ordered.append((best_idx, natural_key(p.name), p))
        else:
            rest.append(p)

    ordered.sort(key=lambda t: (t[0], t[1]))
    rest.sort(key=lambda p: natural_key(p.name))
    unmatched = [codes[i] for i, _ in lowered if i not in used_codes]
    return [t[2] for t in ordered] + rest, unmatched, rest


def read_workbook(path: Path, skip_top: int) -> Dict[str, "pd.DataFrame"]:
    """Reads ALL sheets of a file with a single open (returns {name: df})."""
    ext = path.suffix.lower()
    if ext == ".csv":
        try:
            return {"CSV_Data": read_csv_robust(path, skiprows=skip_top, header=0)}
        except pd.errors.EmptyDataError:
            return {}

    result: Dict[str, pd.DataFrame] = {}
    with pd.ExcelFile(path, engine=excel_engine(ext)) as xls:
        for name in xls.sheet_names:
            result[str(name)] = xls.parse(name, skiprows=skip_top, header=0)
    return result


def dedupe_columns(cols: List[str]) -> List[str]:
    seen: Dict[str, int] = {}
    out: List[str] = []
    for c in cols:
        if c in seen:
            seen[c] += 1
            out.append(f"{c}.{seen[c]}")
        else:
            seen[c] = 0
            out.append(c)
    return out


def clean_frame(df: "pd.DataFrame", skip_bottom: int) -> "pd.DataFrame":
    """Cleans a frame: empty rows/columns, headers, footer."""
    df = df.dropna(how="all")
    empty_unnamed = [
        c for c in df.columns if str(c).startswith("Unnamed:") and df[c].isna().all()
    ]
    if empty_unnamed:
        df = df.drop(columns=empty_unnamed)

    # Headers: strip extra spaces/newlines so that "Price " == "Price"
    df.columns = dedupe_columns([re.sub(r"\s+", " ", str(c)).strip() for c in df.columns])

    # Footer is counted after empty rows have been removed
    if skip_bottom > 0:
        df = df.iloc[:-skip_bottom] if skip_bottom < len(df) else df.iloc[0:0]
    return df.reset_index(drop=True)


def make_safe_sheet_names(names: List[str]) -> List[str]:
    """Sheet names: no forbidden characters, max 31 chars, unique (case-insensitive)."""
    used: set = set()
    result: List[str] = []
    for raw in names:
        base = _BAD_SHEET_CHARS.sub("_", raw).strip("'").strip()[:31] or "Sheet"
        candidate, i = base, 2
        while candidate.casefold() in used:
            suffix = f"~{i}"
            candidate = base[: 31 - len(suffix)] + suffix
            i += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return result


def sanitize_for_excel(df: "pd.DataFrame") -> "pd.DataFrame":
    """Removes control characters that make openpyxl raise IllegalCharacterError."""
    obj_cols = [c for c in df.columns if df[c].dtype.kind == "O"]
    if not obj_cols:
        return df
    df = df.copy()
    for c in obj_cols:
        df[c] = df[c].map(lambda v: _ILLEGAL_XML.sub("", v) if isinstance(v, str) else v)
    return df


def format_sheet(ws, df: "pd.DataFrame") -> None:
    """Bold header, frozen first row, autofilter, column widths."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if df.shape[1] == 0:
        return
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    bold = Font(bold=True)
    for idx, col in enumerate(df.columns, start=1):
        ws.cell(row=1, column=idx).font = bold
        sample = df[col].head(500).astype(str)
        longest = max([len(str(col))] + sample.map(len).tolist())
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 50)


def merge_folder(
    opts: MergeOptions,
    progress: Optional[ProgressCb] = None,
    cancel: Optional[threading.Event] = None,
) -> MergeResult:
    """Main routine: read, clean, merge and save. Raises ValueError on user errors."""
    report: ProgressCb = progress or (lambda done, total, msg: None)
    res = MergeResult()

    if not opts.input_dir.is_dir():
        raise ValueError("Input folder does not exist.")
    if not opts.output_dir.is_dir():
        raise ValueError("Output folder does not exist.")

    # --- input files (skip Excel temp files and our own previous outputs) ---
    files = [
        p
        for p in opts.input_dir.iterdir()
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_EXTS
        and not p.name.startswith(("~$", OUTPUT_PREFIX))
    ]
    if not files:
        raise ValueError("No Excel/CSV files found in the input folder.")

    # --- ordering ---
    codes: List[str] = []
    if opts.order_file:
        try:
            codes = read_order_list(opts.order_file)
            log.info("Order list loaded: %d entries", len(codes))
        except Exception as exc:  # noqa: BLE001
            log.warning("Order file: %s", exc)
            res.warnings.append(f"Could not load the order file: {exc}")

    if codes:
        files, unmatched, unordered = sort_files(files, codes)
        if unmatched:
            res.warnings.append(
                f"Codes from the order file without a matching file ({len(unmatched)}): "
                + ", ".join(unmatched[:10])
                + (" …" if len(unmatched) > 10 else "")
            )
        if unordered:
            res.warnings.append(
                f"Files not in the order list (appended at the end): {len(unordered)}"
            )
    else:
        files.sort(key=lambda p: natural_key(p.name))

    res.files_total = len(files)
    total_steps = len(files) + 1

    # --- reading: sheet key (casefold) -> (display name, list of frames) ---
    sheets: Dict[str, Tuple[str, List[pd.DataFrame]]] = {}
    for n, path in enumerate(files):
        if cancel is not None and cancel.is_set():
            res.cancelled = True
            return res
        report(n, total_steps, f"Loading {n + 1}/{len(files)}: {path.name}")
        try:
            workbook = read_workbook(path, opts.skip_top)
            got_data = False
            for sheet_name, raw in workbook.items():
                df = clean_frame(raw, opts.skip_bottom)
                if df.empty:
                    continue
                if opts.add_source_column and SOURCE_COL not in df.columns:
                    df.insert(0, SOURCE_COL, path.name)
                key = sheet_name.strip().casefold()
                sheets.setdefault(key, (sheet_name.strip(), []))[1].append(df)
                got_data = True
            res.files_ok += 1
            if not got_data:
                res.warnings.append(f"{path.name}: no data left after applying row filters")
        except Exception as exc:  # noqa: BLE001
            log.error("Error in file %s:\n%s", path.name, traceback.format_exc())
            res.errors.append(f"{path.name}: {type(exc).__name__}: {exc}")

    if not sheets:
        raise ValueError("No tabular data left after applying row filters.")

    # --- concatenation (+ splitting when a sheet exceeds the Excel limit) ---
    parts: List[Tuple[str, pd.DataFrame]] = []
    for display, frames in sheets.values():
        merged = pd.concat(frames, ignore_index=True, sort=False)
        if len(merged) > EXCEL_MAX_DATA_ROWS:
            for k, start in enumerate(range(0, len(merged), EXCEL_MAX_DATA_ROWS), start=1):
                parts.append((f"{display} ({k})", merged.iloc[start : start + EXCEL_MAX_DATA_ROWS]))
            res.warnings.append(
                f"Sheet '{display}' exceeds the Excel row limit - split into parts."
            )
        else:
            parts.append((display, merged))

    safe_names = make_safe_sheet_names([name for name, _ in parts])

    # --- saving: temp file first, then atomic replace ---
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = opts.output_dir / f"{OUTPUT_PREFIX}{stamp}.xlsx"
    tmp_file = output_file.with_name(output_file.stem + ".tmp.xlsx")
    report(len(files), total_steps, "Saving output file…")
    try:
        with pd.ExcelWriter(tmp_file, engine="openpyxl") as writer:
            for (_, df), name in zip(parts, safe_names):
                df = sanitize_for_excel(df)
                df.to_excel(writer, sheet_name=name, index=False)
                format_sheet(writer.sheets[name], df)
        os.replace(tmp_file, output_file)
    except Exception:
        try:
            tmp_file.unlink(missing_ok=True)
        except OSError:
            pass
        raise

    res.output_file = output_file
    res.sheets = len(parts)
    res.rows = sum(len(df) for _, df in parts)
    report(total_steps, total_steps, "Done")
    log.info(
        "Merged %d/%d files -> %s (%d sheets, %d rows)",
        res.files_ok, res.files_total, output_file, res.sheets, res.rows,
    )
    return res


# =============================================================================
# 3. GUI
# =============================================================================
class MergerApp(tk.Tk):
    BG = "#F5F5F7"
    FG = "#1D1D1F"

    def __init__(self) -> None:
        super().__init__()
        self.title("Multi-Sheet Excel Merger Studio")
        self.geometry("700x640")
        self.minsize(640, 600)
        self.configure(bg=self.BG)

        self._events: "queue.Queue[tuple]" = queue.Queue()
        self._cancel_event = threading.Event()
        self._worker: Optional[threading.Thread] = None

        self._setup_styles()
        self._build_ui()
        self._load_settings()

        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.after(100, self._poll_events)
        log.info("GUI initialized.")

    # ---- styles --------------------------------------------------------------
    def _setup_styles(self) -> None:
        st = ttk.Style(self)
        st.theme_use("clam")
        st.configure("TLabel", background=self.BG, foreground=self.FG, font=("Segoe UI", 9))
        st.configure("TCheckbutton", background=self.BG, foreground=self.FG, font=("Segoe UI", 9))
        st.configure("TEntry", fieldbackground="#FFFFFF", padding=4)
        st.configure("TSpinbox", fieldbackground="#FFFFFF", padding=3)
        st.configure("Browse.TButton", font=("Segoe UI", 9), background="#0071E3",
                     foreground="white", borderwidth=0, focuscolor="none", padding=(10, 4))
        st.map("Browse.TButton",
               background=[("active", "#0077ED"), ("disabled", "#D2D2D7")],
               foreground=[("disabled", "#86868B")])
        st.configure("Run.TButton", font=("Segoe UI", 10, "bold"), background="#28CD41",
                     foreground="white", borderwidth=0, focuscolor="none", padding=(14, 8))
        st.map("Run.TButton",
               background=[("active", "#24B338"), ("disabled", "#D2D2D7")],
               foreground=[("disabled", "#86868B")])
        st.configure("Cancel.TButton", font=("Segoe UI", 10), background="#FF5F56",
                     foreground="white", borderwidth=0, focuscolor="none", padding=(14, 8))
        st.map("Cancel.TButton",
               background=[("active", "#E0443E"), ("disabled", "#D2D2D7")],
               foreground=[("disabled", "#86868B")])
        st.configure("Horizontal.TProgressbar", troughcolor="#E8E8ED", background="#0071E3")

    # ---- UI construction -------------------------------------------------------
    def _build_ui(self) -> None:
        self.var_input = tk.StringVar()
        self.var_order = tk.StringVar()
        self.var_output = tk.StringVar()
        self.var_top = tk.StringVar(value="0")
        self.var_bottom = tk.StringVar(value="0")
        self.var_source = tk.BooleanVar(value=False)
        self.var_open = tk.BooleanVar(value=True)
        self._persisted = {
            "input": self.var_input, "order": self.var_order, "output": self.var_output,
            "top": self.var_top, "bottom": self.var_bottom,
            "source": self.var_source, "open": self.var_open,
        }

        content = tk.Frame(self, bg=self.BG)
        content.pack(fill="both", expand=True, padx=25, pady=(12, 4))
        content.columnconfigure(0, weight=1)

        self._path_row(content, 0, "1. Input folder (files to merge):",
                       self.var_input, lambda: self._browse_dir(self.var_input, "Select input folder"))
        self._path_row(content, 1, "2. Order file (optional, column A):",
                       self.var_order, self._browse_order)
        self._path_row(content, 2, "3. Output folder for the merged file:",
                       self.var_output, lambda: self._browse_dir(self.var_output, "Select output folder"))

        opts = tk.LabelFrame(content, text=" Row cleanup options (applied to every sheet) ",
                             bg=self.BG, fg=self.FG, font=("Segoe UI", 9, "bold"), padx=10, pady=8)
        opts.grid(row=6, column=0, columnspan=2, sticky="ew", pady=(12, 4))
        opts.columnconfigure(0, weight=1)
        opts.columnconfigure(1, weight=1)
        vcmd = (self.register(lambda v: v == "" or v.isdigit()), "%P")

        ttk.Label(opts, text="Skip rows at the top (before the header):").grid(row=0, column=0, sticky="w")
        ttk.Spinbox(opts, from_=0, to=9999, width=10, textvariable=self.var_top,
                    validate="key", validatecommand=vcmd).grid(row=1, column=0, sticky="ew", padx=(0, 10), pady=(2, 0))
        ttk.Label(opts, text="Skip rows at the bottom (footer / totals):").grid(row=0, column=1, sticky="w")
        ttk.Spinbox(opts, from_=0, to=9999, width=10, textvariable=self.var_bottom,
                    validate="key", validatecommand=vcmd).grid(row=1, column=1, sticky="ew", pady=(2, 0))
        ttk.Label(opts, text="Empty rows are ignored when counting footer rows.",
                  foreground="#86868B").grid(row=2, column=0, columnspan=2, sticky="w", pady=(6, 0))

        flags = tk.Frame(content, bg=self.BG)
        flags.grid(row=7, column=0, columnspan=2, sticky="w", pady=(8, 0))
        ttk.Checkbutton(flags, text="Add a column with the source file name (_Source_File)",
                        variable=self.var_source).pack(anchor="w")
        ttk.Checkbutton(flags, text="Open the output folder when finished",
                        variable=self.var_open).pack(anchor="w")

        buttons = tk.Frame(self, bg=self.BG)
        buttons.pack(pady=(8, 6))
        self.btn_run = ttk.Button(buttons, text="RUN MERGE (SHEET BY SHEET)",
                                  style="Run.TButton", command=self._start)
        self.btn_run.pack(side="left", padx=4)
        self.btn_cancel = ttk.Button(buttons, text="Cancel", style="Cancel.TButton",
                                     command=self._cancel, state=tk.DISABLED)
        self.btn_cancel.pack(side="left", padx=4)

        self.progress = ttk.Progressbar(self, mode="determinate", maximum=100)
        self.progress.pack(fill="x", padx=25, pady=(4, 4))
        self.lbl_status = tk.Label(self, text="Ready to merge", bg=self.BG, fg="#86868B",
                                   font=("Segoe UI", 9, "bold"), wraplength=620, justify="center")
        self.lbl_status.pack(pady=(0, 10))

    def _path_row(self, parent: tk.Frame, idx: int, label: str, var: tk.StringVar, command) -> None:
        ttk.Label(parent, text=label).grid(row=idx * 2, column=0, columnspan=2, sticky="w", pady=(8, 2))
        ttk.Entry(parent, textvariable=var).grid(row=idx * 2 + 1, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(parent, text="Browse…", style="Browse.TButton", command=command).grid(row=idx * 2 + 1, column=1)

    # ---- browsing ---------------------------------------------------------------
    def _browse_dir(self, var: tk.StringVar, title: str) -> None:
        path = filedialog.askdirectory(title=title, initialdir=var.get() or None)
        if path:
            var.set(str(Path(path)))

    def _browse_order(self) -> None:
        path = filedialog.askopenfilename(
            title="Select order file",
            filetypes=[("Excel / CSV files", "*.xlsx *.xlsm *.xlsb *.xls *.csv"), ("All files", "*.*")],
        )
        if path:
            self.var_order.set(str(Path(path)))

    # ---- settings ---------------------------------------------------------------
    def _load_settings(self) -> None:
        try:
            data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return
        for key, var in self._persisted.items():
            if key in data:
                try:
                    var.set(data[key])
                except tk.TclError:
                    pass

    def _save_settings(self) -> None:
        try:
            data = {k: v.get() for k, v in self._persisted.items()}
            SETTINGS_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        except (OSError, tk.TclError):
            pass

    # ---- processing control ---------------------------------------------------------
    def _set_running(self, running: bool) -> None:
        self.btn_run.config(state=tk.DISABLED if running else tk.NORMAL)
        self.btn_cancel.config(state=tk.NORMAL if running else tk.DISABLED)

    def _set_status(self, text: str, color: str) -> None:
        self.lbl_status.config(text=text, fg=color)

    @staticmethod
    def _to_int(value: str) -> int:
        return int(value) if value.strip().isdigit() else 0

    def _start(self) -> None:
        input_dir = self.var_input.get().strip()
        output_dir = self.var_output.get().strip()
        if not input_dir or not output_dir:
            self._set_status("Error: input and output folders are required.", "#D70015")
            return

        order = self.var_order.get().strip()
        opts = MergeOptions(
            input_dir=Path(input_dir),
            output_dir=Path(output_dir),
            order_file=Path(order) if order else None,
            skip_top=self._to_int(self.var_top.get()),
            skip_bottom=self._to_int(self.var_bottom.get()),
            add_source_column=self.var_source.get(),
        )
        self._save_settings()
        self._cancel_event.clear()
        self._set_running(True)
        self.progress.config(value=0)
        self._set_status("Processing… please wait.", "#0071E3")
        self._worker = threading.Thread(target=self._worker_main, args=(opts,), daemon=True)
        self._worker.start()

    def _cancel(self) -> None:
        self._cancel_event.set()
        self._set_status("Cancelling…", "#86868B")

    def _worker_main(self, opts: MergeOptions) -> None:
        """Runs in a worker thread - talks to the GUI ONLY through the queue."""
        try:
            result = merge_folder(
                opts,
                progress=lambda d, t, m: self._events.put(("progress", d, t, m)),
                cancel=self._cancel_event,
            )
            self._events.put(("done", result))
        except PermissionError:
            self._events.put(("error", "Write access denied. Close the file in Excel or choose another folder."))
        except Exception as exc:  # noqa: BLE001
            log.error("Merge failed:\n%s", traceback.format_exc())
            self._events.put(("error", f"{type(exc).__name__}: {exc}"))

    def _poll_events(self) -> None:
        try:
            while True:
                event = self._events.get_nowait()
                kind = event[0]
                if kind == "progress":
                    _, done, total, msg = event
                    self.progress.config(value=100 * done / max(total, 1))
                    self._set_status(msg, "#0071E3")
                elif kind == "done":
                    self._on_done(event[1])
                elif kind == "error":
                    self._set_running(False)
                    self.progress.config(value=0)
                    self._set_status(f"Error: {event[1]}", "#D70015")
        except queue.Empty:
            pass
        self.after(100, self._poll_events)

    def _on_done(self, res: MergeResult) -> None:
        self._set_running(False)
        if res.cancelled:
            self.progress.config(value=0)
            self._set_status("Cancelled - no file was saved.", "#86868B")
            return

        self.progress.config(value=100)
        has_issues = bool(res.errors or res.warnings)
        out_name = res.output_file.name if res.output_file else ""
        msg = (f"Done: {res.files_ok}/{res.files_total} files, {res.sheets} sheets, "
               f"{res.rows:,} rows → {out_name}")
        self._set_status(msg, "#FF9500" if res.errors else "#28CD41")

        if has_issues:
            messagebox.showwarning("Merge report", self._format_report(res))
        if self.var_open.get() and res.output_file and hasattr(os, "startfile"):
            try:
                os.startfile(res.output_file.parent)  # type: ignore[attr-defined]
            except OSError:
                pass

    @staticmethod
    def _format_report(res: MergeResult, limit: int = 15) -> str:
        lines: List[str] = []
        if res.errors:
            lines.append(f"Files with errors ({len(res.errors)}):")
            lines += [f"  • {e}" for e in res.errors[:limit]]
            if len(res.errors) > limit:
                lines.append("  …")
        if res.warnings:
            lines.append(f"\nWarnings ({len(res.warnings)}):")
            lines += [f"  • {w}" for w in res.warnings[:limit]]
            if len(res.warnings) > limit:
                lines.append("  …")
        lines.append(f"\nFull log: {LOG_FILE}")
        return "\n".join(lines)

    def _on_close(self) -> None:
        self._save_settings()
        self._cancel_event.set()
        self.destroy()


def main() -> None:
    try:
        enable_dpi_awareness()
        app = MergerApp()
        log.info("Main loop started.")
        app.mainloop()
    except Exception as exc:  # noqa: BLE001
        log.critical("FATAL CRASH:\n%s", traceback.format_exc())
        show_native_error("Critical application error", f"An error occurred:\n\n{exc}\n\nSee: {LOG_FILE}")
        sys.exit(1)


if __name__ == "__main__":
    main()
